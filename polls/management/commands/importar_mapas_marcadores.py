"""
Comando para importar o Mapa de Marcadores da planilha Excel para o banco de dados.

Uso:
    python manage.py importar_mapas_marcadores "caminho/para/Mapa de Marcadores.xlsx"
    python manage.py importar_mapas_marcadores "caminho/para/Mapa.xlsx" --dry-run
"""

import re
from decimal import Decimal, InvalidOperation
from django.core.management.base import BaseCommand, CommandError
from openpyxl import load_workbook

from polls.models import MapaMarcador


# Mapeamento de texto para nível
QUAL_PARA_NIVEL = {
    'exito': 3,
    'intermediario': 2,
    'ausencia': 1,
    'sucesso': 3,
    'parcial': 2,
    'insucesso': 1,
    'conservador': 3,
    'intermediario': 2,
    'nao_conservador': 1,
}


def _norm(valor):
    """Normaliza um valor para string, tratando None e removendo espaços extras"""
    if valor is None:
        return ''
    if isinstance(valor, (int, float, Decimal)):
        return str(valor).strip()
    return str(valor).strip()


def _slugify(texto):
    """Cria um slug simples a partir do texto"""
    if not texto:
        return ''
    texto = texto.lower()
    texto = re.sub(r'[^a-z0-9_]', '_', texto)
    texto = re.sub(r'_+', '_', texto)
    return texto.strip('_')


def _eh_token_nivel(texto):
    """Verifica se o texto é um token de nível (1, 2, 3, Exito, Intermediario, etc)"""
    if not texto:
        return False
    texto = texto.lower().strip()
    # Números
    if texto in ['1', '2', '3']:
        return True
    # Palavras-chave
    if texto in ['exito', 'intermediario', 'ausencia', 'sucesso', 'parcial', 'insucesso']:
        return True
    # Padrão "Nível X"
    if re.match(r'n[ií]vel\s*[123]', texto):
        return True
    return False


def _nivel_de_token(texto):
    """Extrai o nível de um token textual"""
    if not texto:
        return None
    texto = texto.lower().strip()
    
    # Número direto
    if texto in ['1', '2', '3']:
        return int(texto)
    
    # Padrão "Nível X"
    match = re.search(r'n[ií]vel\s*([123])', texto)
    if match:
        return int(match.group(1))
    
    # Mapeamento de palavras
    return QUAL_PARA_NIVEL.get(texto)


def _parse_num(valor):
    """Converte valor para Decimal, tratando None e strings vazias"""
    if valor is None:
        return None
    if isinstance(valor, (int, float)):
        return Decimal(str(valor))
    if isinstance(valor, Decimal):
        return valor
    texto = str(valor).strip()
    if not texto:
        return None
    try:
        return Decimal(texto)
    except (InvalidOperation, ValueError):
        return None


def _pad(lista, tamanho):
    """Garante que a lista tenha o tamanho mínimo"""
    if len(lista) >= tamanho:
        return lista
    return lista + [None] * (tamanho - len(lista))


class Command(BaseCommand):
    help = 'Importa o Mapa de Marcadores da planilha Excel para o banco de dados'

    def add_arguments(self, parser):
        parser.add_argument('caminho', type=str, help='Caminho para o arquivo Excel')
        parser.add_argument('--dry-run', action='store_true', help='Apenas simula a importação')

    def handle(self, *args, **options):
        caminho = options['caminho']
        self.dry_run = options['dry_run']
        self.criados = 0
        self.atualizados = 0

        try:
            wb = load_workbook(caminho, data_only=True)
        except Exception as e:
            raise CommandError(f'Erro ao abrir a planilha: {e}')

        self.stdout.write(f'Processando: {caminho}')
        
        # Configurações por aba
        configs = [
            {
                'jogo': 'piaget',
                'col_tela': 1,
                'col_opcao': 2,
                'col_sigla': 3,
                'col_rotulo': 4,
                'col_valor': 5,
                'col_descricao': 6,
                'col_nivel': 7,
                'tipo': 'nivel',
            },
            {
                'jogo': 'volta_casa',
                'col_tela': 0,
                'col_opcao': 1,
                'dim_cols': [3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16, 17],
                'indiv_cols': [],
                'tipo': 'dimensional',
            },
            {
                'jogo': 'volta_casa_cog',
                'col_tela': 0,
                'col_opcao': 1,
                'dim_cols': [3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16, 17, 18, 19],
                'indiv_cols': [],
                'tipo': 'dimensional',
            },
        ]

        # Mapear abas por nome
        abas_esperadas = {
            'piaget': 'piaget',
            'volta_casa': 'volta_casa',
            'volta_casa_cog': 'volta_casa_cog',
        }
        
        for cfg in configs:
            jogo = cfg['jogo']
            nome_aba = abas_esperadas.get(jogo)
            if not nome_aba:
                continue
                
            # Encontrar a aba correspondente
            aba_encontrada = None
            for nome in wb.sheetnames:
                if nome.lower() == nome_aba.lower():
                    aba_encontrada = nome
                    break
            
            if not aba_encontrada:
                self.stdout.write(self.style.WARNING(f'Aba "{nome_aba}" não encontrada'))
                continue
            
            sheet = wb[aba_encontrada]
            rows = list(sheet.iter_rows(values_only=True))
            
            # Pular cabeçalho
            if rows:
                rows = rows[1:]
            
            self.stdout.write(f'Processando aba: {aba_encontrada} ({len(rows)} linhas)')
            
            if cfg['tipo'] == 'nivel':
                self._importar_nivel(cfg, rows)
            else:
                self._importar_dimensional(cfg, rows)

        self.stdout.write(self.style.SUCCESS(
            f'\nConcluído. Criados: {self.criados} | Atualizados: {self.atualizados}'
        ))
        if self.dry_run:
            self.stdout.write(self.style.WARNING('dry-run: nada foi gravado'))

    def _importar_nivel(self, cfg, rows):
        """Importa jogos por nível (Piaget)"""
        tamanho = max(cfg['col_valor'], cfg['col_descricao']) + 1
        fase_atual = ''
        vistos = set()
        
        for row in rows:
            row = _pad(row, tamanho)
            tela = _norm(row[cfg['col_tela']])
            if tela:
                fase_atual = tela
            
            if not fase_atual:
                continue
            
            opcao = _norm(row[cfg['col_opcao']])
            if not opcao:
                continue
            
            # Extrair nível de diferentes colunas
            cand1 = _norm(row[cfg['col_rotulo']])
            cand2 = _norm(row[cfg['col_nivel']])
            
            nivel = _nivel_de_token(cand1) or _nivel_de_token(cand2)
            rotulo = ''
            
            for c in (cand1, cand2):
                if c and not _eh_token_nivel(c):
                    rotulo = c
                    break
            
            if nivel is None and rotulo:
                nivel = QUAL_PARA_NIVEL.get(_slugify(rotulo))
            
            if nivel is None and not rotulo:
                continue
            
            # Gerar chave única
            base = _slugify(rotulo) or _slugify(opcao) or (f'nivel_{nivel}' if nivel else 'opcao')
            chave = self._chave_unica(base, fase_atual, vistos)
            
            valor = _parse_num(row[cfg['col_valor']]) or Decimal('0')
            
            self._gravar(cfg['jogo'], fase_atual, chave, {
                'sigla': _norm(row[cfg['col_sigla']])[:3],
                'nivel': nivel,
                'rotulo': rotulo or opcao,
                'valor_quantitativo': valor,
                'descricao_qualitativa': _norm(row[cfg['col_descricao']]) or None,
            })

    def _importar_dimensional(self, cfg, rows):
        """Importa jogos dimensionais (Volta pra Casa e Cognitivo)"""
        tamanho = max(cfg.get('dim_cols', []) + cfg.get('indiv_cols', []) + [cfg['col_opcao']]) + 1
        fase_atual = ''
        vistos = set()
        bloco = None
        
        def flush(b):
            if not b or not b['opcao'] or not b['fase']:
                return
            chave = self._chave_unica(_slugify(b['opcao'] or 'opcao'), b['fase'], vistos)
            self._gravar(cfg['jogo'], b['fase'], chave, {
                'sigla': '',
                'nivel': None,
                'rotulo': b['opcao'],
                'valor_quantitativo': b['valor'],
                'descricao_qualitativa': ' '.join(b['dims'])[:5000] or None,
            })
        
        for row in rows:
            row = _pad(row, tamanho)
            tela = _norm(row[cfg['col_tela']])
            if tela:
                fase_atual = tela
            
            if not fase_atual:
                continue
            
            opcao = _norm(row[cfg['col_opcao']])
            # Pular linha de "Possibilidades"
            if opcao and opcao.lower().startswith('possibilidades'):
                continue
            
            if opcao:
                flush(bloco)
                bloco = {
                    'fase': fase_atual,
                    'opcao': opcao,
                    'valor': Decimal('0'),
                    'dims': []
                }
            
            if bloco is None:
                continue
            
            # Somar valores individuais
            for c in cfg.get('indiv_cols', []):
                if c < len(row):
                    num = _parse_num(row[c])
                    if num is not None:
                        bloco['valor'] += num
            
            # Coletar dimensões
            for c in cfg.get('dim_cols', []):
                if c < len(row):
                    nome = _norm(row[c])
                    if nome and nome != '-':
                        bloco['dims'].append(nome)
        
        flush(bloco)

    def _chave_unica(self, base, fase, vistos):
        """Gera uma chave única para a combinação fase + chave"""
        chave = base
        suffix = 2
        while (fase, chave) in vistos:
            chave = f'{base}_{suffix}'
            suffix += 1
        vistos.add((fase, chave))
        return chave

    def _gravar(self, jogo, fase, chave, dados):
        """Grava ou atualiza um registro no MapaMarcador"""
        if self.dry_run:
            nivel_str = f"N={dados['nivel']}" if dados['nivel'] is not None else "N=None"
            self.stdout.write(f'  [dry] {jogo}/{fase}/{chave} | {nivel_str} | valor={dados["valor_quantitativo"]}')
            self.criados += 1
            return
        
        defaults = {
            'sigla': dados.get('sigla', ''),
            'nivel': dados.get('nivel'),
            'rotulo': dados.get('rotulo', ''),
            'valor_quantitativo': dados['valor_quantitativo'],
            'descricao_qualitativa': dados.get('descricao_qualitativa'),
        }
        
        obj, created = MapaMarcador.objects.update_or_create(
            jogo=jogo,
            fase=fase,
            resposta_chave=chave,
            defaults=defaults
        )
        
        if created:
            self.criados += 1
        else:
            self.atualizados += 1